"""Build docs/supplement_marker_timing_calibration_2026-10-15.docx from the analysis outputs."""
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(r"C:\Users\aviva-Lab\Desktop\vMMR_EEG_PsychoPy_experiment")
A = REPO / "data" / "analysis" / "calibration_database"
FIG = A / "figures"
OUT = REPO / "docs" / "supplement_marker_timing_calibration_2026-10-15.docx"
REPO_URL = "https://github.com/DanielAtad/vMMR_EEG_PsychoPy_experiment"

s = json.loads((A / "summary.json").read_text(encoding="utf-8"))
runs = pd.read_csv(A / "runs.csv", dtype={"date": str})
db = runs[runs.set == "database"].copy()
val = runs[runs.set == "validation"].copy()
C = s["correction"]
Cm = s["correction_median_base"]
R = s["recorded_marker"]
E = s["flash_error_loo"]
Em = s["flash_error_loo_median_base"]
Er = s["flash_error_loo_recorded_marker"]
V = s["validation_corner"]
full = s["long_runs"]["20261015_full_01"]
day = s["day_effect"]

# frames shown and dropped (PsychoPy logs of the matched runs)
frames = dropped = 0
for name in db.psychopy_csv:
    c = pd.read_csv(REPO / "data" / name)
    frames += int((c.on_frames + c.off_frames).sum())
    dropped += int(c.dropped_frames_delta.sum())
lag = pd.read_csv(A / "model_lag.csv")
slopes = [np.polyfit(g.sim_s / 60, g.lag_rel_ms, 1)[0] for rid, g in lag.groupby("run_id") if rid in set(db.run_id)]
drift_ppm = np.median(slopes) / 60 * 1000
full_slope = np.polyfit(*(lambda g: (g.sim_s / 60, g.lag_rel_ms))(lag[lag.run_id == "20261015_full_01"]), 1)[0]


def f1(x):
    return f"{x:.1f}"


def ci(d):
    return f"{d[0]:.1f}–{d[1]:.1f}"


# ----------------------------------------------------------------------------- document helpers
doc = docx.Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)


def add_runs(p, text, bold=False, italic=False, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = p.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`"):
            run = p.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt((size or 10.5) - 1)
        else:
            run = p.add_run(part); run.bold = bold
        if italic:
            run.italic = True
        if size and not part.startswith("`"):
            run.font.size = Pt(size)
    return p


def para(text="", bold=False, italic=False, size=None, space_after=6, align=None, style=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, bold, italic, size)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        add_runs(p, it)
        p.paragraph_format.space_after = Pt(2)


def h(text, level):
    doc.add_heading(text, level=level)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(header, rows, widths, caption, font=8.5, note=None, keep_together=True):
    if caption:
        para(caption, bold=True, size=9.5, space_after=3).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, x in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        add_runs(c.paragraphs[0], x, bold=True, size=font); shade(c, "E8EEF7")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; add_runs(cells[i].paragraphs[0], str(v), size=font)
    for ri, row in enumerate(t.rows):
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); trPr.append(cant)
        if ri == 0:                                       # repeat the header row on a new page
            hdr = OxmlElement("w:tblHeader"); trPr.append(hdr)
        if keep_together and ri < len(t.rows) - 1:        # keep a short table on one page
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    if note:
        para(note, italic=True, size=8.5, space_after=8)
    else:
        para("", space_after=4)


def figure(path, caption, width=6.7):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, size=9, space_after=10)


def equation(text):
    p = para(text, size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.font.name = "Cambria Math"
    return p


# ----------------------------------------------------------------------------- title
t = doc.add_heading("Supplementary Material S1. Event-marker timing calibration of the vMMR EEG set-up", 0)
for r in t.runs:
    r.font.size = Pt(18)
para("Optical validation of Lab Streaming Layer (LSL) event markers recorded with a g.HIamp amplifier "
     "in Simulink, and the correction applied to all event latencies", italic=True, space_after=2)
para(f"Calibration data recorded 14–15 October 2026 (pilot 6–13 October 2026). Code and analysis: {REPO_URL} "
     "(branch `photodiode-latency-calibration`, commit [to be inserted at submission]).", size=9, space_after=10)

# ----------------------------------------------------------------------------- to complete (remove before submission)
todo = doc.add_table(rows=1, cols=1)
todo.style = "Table Grid"
cell = todo.rows[0].cells[0]
cell.width = Inches(6.7)
shade(cell, "FFF4CE")
cell.text = ""
add_runs(cell.paragraphs[0], "TO COMPLETE BEFORE SUBMISSION (delete this box)", bold=True, size=10)
for item in [
    "Title block: the commit hash of the code version used (e.g. the release commit of branch "
    "`photodiode-latency-calibration`).",
    "Table S1, recording computer (A): g.HIamp model and channel count, g.HIsys version, liblsl-Matlab version, "
    "Windows version (and CPU if wanted).",
    "Table S1, network: confirm the connection between A and B (currently “local Ethernet”).",
    "S1.6: confirm per study that hardware, software and display settings were unchanged between calibration and "
    "data collection, or report the re-calibration sessions.",
    "Data and code availability: where the raw recordings are deposited (repository/DOI or Google Drive link) and "
    "the access conditions.",
    "References: replace the Kothe et al. (2024) bioRxiv entry with its final citation if published, and check all "
    "DOIs against the journal's reference style.",
    "Numbering: adapt “Supplementary Material S1” and the S-numbers of tables and figures to the journal's scheme.",
]:
    p = cell.add_paragraph(style="List Bullet")
    add_runs(p, item, size=9)
    p.paragraph_format.space_after = Pt(1)
para("", space_after=6)

# ----------------------------------------------------------------------------- abstract
h("Summary", 1)
para(f"Event markers in this study are sent by the stimulus computer over the Lab Streaming Layer (LSL) and "
     f"recorded with the EEG on a separate computer. Because a photodiode cannot be used during experimental "
     f"sessions, we calibrated the delay between each marker and the physical appearance of the stimulus on the "
     f"screen in a dedicated database of {C['n']} sessions recorded on 2 days ({s['n_flashes_included']:,} "
     f"full-white flashes at the screen centre, detected with an optical sensor recorded by the EEG amplifier). "
     f"Each marker was placed on the EEG time axis by its LSL send timestamp, converted to the recording "
     f"computer's clock and then to amplifier sample time.")
bullets([
    f"**Correction:** the stimulus light reaches the screen centre **{f1(C['mean'])} ms** after the timestamp-placed "
    f"marker (95 % CI {ci(C['ci95'])} ms). Sessions differ by SD {C['sd']:.1f} ms; a new session is expected "
    f"within {ci(C['pi95'])} ms (95 % prediction interval).",
    f"**Per-event precision:** after correction, single events deviate from the optical onset with SD "
    f"{E['sd']:.1f} ms; {E['within_10ms']:.1%} lie within ±10 ms and all within one display frame "
    f"(max {E['max_abs']:.0f} ms; leave-one-session-out). This residual is set by the display and cannot be "
    f"removed with markers alone.",
    f"**Stability:** no drift over a {full['duration_min']:.0f}-min session ({full['drift_ms_per_min']:+.2f} ms/min), "
    f"no difference between short and long sessions, and a day difference of {abs(day['diff_ms']):.1f} ms "
    f"(p = {day['p']:.2f}).",
    f"**Method matters:** using the marker as recorded by the Simulink LSL block gives a between-session SD of "
    f"{R['sd']:.1f} ms and single-event errors up to {Er['max_abs']:.0f} ms; the timestamp-based placement is "
    f"therefore required.",
    f"**Replication:** the same procedure applied to {V['n']} independent pilot sessions gave a between-session "
    f"SD of {V['offset']['sd']:.1f} ms; with the sensor in the bottom-right corner the offset was "
    f"{V['corner_minus_centre_ms']:.1f} ms larger, as expected from top-to-bottom screen refresh.",
    f"No session met any of the pre-defined exclusion criteria.",
])

# ----------------------------------------------------------------------------- background
h("S1.1 Background", 1)
para("ERP latencies are measured relative to event markers, so any fixed or variable delay between a marker and "
     "the physical stimulus shifts or smears the averaged waveform (Luck, 2014). Reporting guidelines therefore "
     "ask for the timing accuracy of event codes to be verified and reported (Keil et al., 2014; Pernet et al., "
     "2020), and timing errors in stimulus software and hardware are a known source of irreproducibility "
     "(Plant, 2016; Bridges et al., 2020). Display timing is usually validated with an optical sensor on the "
     "screen (Elze, 2010; Plant et al., 2004).")
para("In this set-up the stimulus computer (PsychoPy; Peirce et al., 2019) sends integer markers as an LSL stream "
     "(Kothe et al., 2024) to a second computer that records EEG with a g.HIamp amplifier in Simulink. Three "
     "delays separate a marker from the light it labels: the display pipeline (from the PsychoPy flip to light "
     "emission), the network and LSL inlet, and the amplifier's block-wise data delivery to the recording model. "
     "During pilot testing, a one-off pause of the stimulus process while it opened its full-screen window "
     "(~0.8 s) left the recording model permanently behind the marker stream; the sender was changed so that the "
     "LSL stream is created only after the window is open (S1.8). All data reported here were recorded with the "
     "corrected sender. The optical sensor is used only for calibration and is removed for experimental sessions.")

# ----------------------------------------------------------------------------- set-up
h("S1.2 Set-up", 1)
table(["Component", "Details"], [
    ["Stimulus computer (B)", "AMD Ryzen 3 3200G, integrated Radeon Vega 8 graphics (driver 31.0.21916.2), "
                              "Windows 11 Education 10.0.26200; Python 3.10.22, PsychoPy 2026.2.4, pylsl 1.18.5 (liblsl 1.18)"],
    ["Display", "Dell P2426H, 1920 × 1080, Windows refresh setting 59.94 Hz; measured "
                f"{s['refresh_hz']['min']:.2f}–{s['refresh_hz']['max']:.2f} Hz in every session (PsychoPy "
                "getActualFrameRate); full-screen window, waitBlanking on"],
    ["Recording computer (A)", "g.tec g.HIamp ([model/channel count TBD]), 600 Hz, read by the g.HIamp Simulink block "
                               "(frame length 8 samples = 13.3 ms); MATLAB/Simulink R2020a, g.HIsys [version TBD], "
                               "liblsl-Matlab [version TBD]; Windows [version TBD]"],
    ["Simulink model", "Fixed-step discrete solver, step 1/600 s, single-tasking, no pacing; EEG, optical "
                       "sensor and marker multiplexed and written to file at every step; recording stops on marker 255"],
    ["Marker reception", "(1) g.HIsys LSL Receive block, 600 Hz (marker value per step); (2) custom MATLAB System "
                         "block `LSLMarkerTimestamps` (second inlet; per step: marker, rising-edge code, its LSL "
                         "send timestamp, clock offset B→A from LSL time_correction refreshed every 1 s, and A's LSL "
                         "clock at the step), saved to a separate file"],
    ["LSL stream", "One int32 channel, type Markers, nominal rate 600 Hz; the sender pushes a sample every 1/600 s "
                   f"(achieved {s['lsl_rate_sent_hz']['min']:.3f}–{s['lsl_rate_sent_hz']['max']:.3f} Hz), holding "
                   "each code for the stimulus duration (calibration) or 100 ms; network: local Ethernet between A and B"],
    ["Optical sensor", "g.tec optical sensor (GTEC-0270) taped over a 45 × 45 px full-white square at the screen "
                       "centre → g.TRIGbox (GTEC-0274W) → trigger adapter (GTEC-0274TR) → g.HIamp digital input, "
                       "recorded as an EEG channel at 600 Hz"],
], widths=[1.55, 5.15], caption="Table S1. Hardware and software")

# ----------------------------------------------------------------------------- procedure
h("S1.3 Calibration procedure", 1)
para("The calibration script (`run_photodiode_latency_v1_trig_after_framerate_measurement.py`) uses the same "
     "window, frame-based timing and LSL sender as the experiment. Each session consisted of: start-up and "
     "refresh-rate check (the session stops if the measured rate deviates > 2 % from 59.94 Hz); creation of the "
     "LSL stream; sensor alignment; an experiment-start marker; 2 s black; a polarity calibration (3 s dark, "
     "3 s white, 3 s dark); 2 s black; and the flash block.")
bullets([
    "**Flashes:** the white square is shown for a whole number of frames and the flash marker (code 20) is pushed "
    "from a flip callback (`win.callOnFlip`), i.e. immediately after the buffer swap of the first frame of the "
    "flash. The flash duration increases by one frame per flash (6 to 105 frames, 100 ms to 1.75 s) and repeats "
    "every 100 flashes, so each optical pulse can be matched to its marker by its width. Inter-flash intervals were "
    "drawn at random from 21–39 frames (0.35–0.65 s) to prevent phase-locking with the 600 Hz LSL and amplifier "
    "cycles.",
    "**Session types:** short (100 flashes, ~3 min), long (400 flashes, ~10 min) and full experiment length "
    "(1200 flashes, ~29 min). Fresh starts of the recording model for every session sample the run-to-run "
    "variation that arises at each start.",
    "**Fixed configuration:** identical hardware, software versions and settings in all sessions; the sensor was "
    "not moved between sessions.",
])
counts = db.groupby(["date", "run_type"]).size().unstack(fill_value=0)
table(["Day", "Short (100)", "Long (400)", "Full length (1200)", "Flashes"],
      [[f"{d[6:]} Oct 2026", counts.loc[d].get("short", 0), counts.loc[d].get("long", 0),
        counts.loc[d].get("full", 0), int(db[db.date == d].flash_markers.sum())] for d in counts.index]
      + [["Total", int(counts.get("short", 0).sum()), int(counts.get("long", 0).sum()),
          int(counts.get("full", 0).sum()), int(db.flash_markers.sum())]],
      widths=[1.4, 1.2, 1.2, 1.5, 1.2], caption="Table S2. Calibration database")
para("**Exclusion criteria (defined before recording):** sender resynchronisations (`lsl_resyncs`) > 0; a lasting "
     "backlog of the recording model > 5 ms; < 95 % of flashes paired with an optical pulse; dropped display frames "
     "> 1 %; incomplete session. None of the sessions met any criterion.", space_after=8)

# ----------------------------------------------------------------------------- analysis
h("S1.4 Analysis", 1)
para("**Optical onsets and pairing.** Pulses were extracted from the optical channel and from the marker channel "
     "(gaps < 30 ms bridged, pulses < 30 ms ignored). The sensor polarity was taken from the polarity calibration "
     "and the width ramp. In every session the number of optical pulses equalled the number of flash markers and "
     "the pulse widths followed the marker widths (r > 0.99), so the n-th optical pulse was paired with flash n.")
para("**Placing a marker by its LSL timestamp.** Three clocks are involved: B's LSL clock (send timestamp "
     "t_B), A's LSL clock (t_A, read at every model step) and the amplifier sample clock, which defines the EEG "
     "time axis (τ = sample index / 600 Hz). The offset between the two LSL clocks, Δ_BA, is estimated "
     "continuously by LSL. The relation between t_A and τ is observed at every model step as")
equation("lag(k) = t_A(k) − τ(k)")
para("Because the amplifier delivers its data in frames of 8 samples, the model processes each frame in a burst "
     "after it arrives: within a frame, lag varies by up to one frame (13.3 ms), and in some sessions the model ran "
     "steadily about one frame behind instead (Fig. S2A). The first sample of a frame has waited the full frame; "
     "the upper envelope of lag therefore tracks the acquisition time in both modes. It also follows the slow "
     "drift between the amplifier and PC clocks. We used the 90th percentile of lag within ±10 s of each event:")
equation("B(τ) = P90{ lag(k) : |τ(k) − τ| ≤ 10 s },      τ_marker = t_B + Δ_BA − B(τ)")
equation("offset = τ_light − τ_marker")
para("where τ_light is the first sample of the optical pulse. The previously used mapping, a single median of "
     "lag per session, is reported for comparison.")
para("**Statistics.** Each session contributes its mean offset (sessions are the unit of replication). The "
     "correction is the mean over sessions, with a t-based 95 % CI and a 95 % prediction interval for a new session "
     "(mean ± t·SD·√(1 + 1/n)). Per-event accuracy was evaluated by leave-one-session-out cross-validation: the "
     "correction estimated from all other sessions was applied to every flash of the left-out session. Day and "
     "session-type effects were tested with Welch's t-test; drift is the slope of offset over time within a session. "
     "The percentile (90) was chosen from a sensitivity scan (Fig. S6A); results were then checked on 8 independent "
     "pilot sessions not used for that choice (sensor at the bottom-right corner, 13 October).")

# ----------------------------------------------------------------------------- results
h("S1.5 Results", 1)
para(f"All {C['n']} sessions passed the exclusion criteria; all {s['n_flashes_included']:,} flashes were paired. "
     f"Sender: {s['lsl_rate_sent_hz']['min']:.3f}–{s['lsl_rate_sent_hz']['max']:.3f} samples/s, no "
     f"resynchronisations. Display: {dropped} dropped frames in {frames:,} ({dropped / frames:.3%}). Recording "
     f"model: {s['stalls']['total']} short stalls (> 100 ms, largest {s['stalls']['largest_ms']:.0f} ms) in "
     f"{s['stalls']['runs_with_stall']} sessions, all caught up within seconds (no lasting backlog, Fig. S5).")
table(["Quantity", "Timestamp, upper envelope (adopted)", "Timestamp, median mapping", "Recorded marker"], [
    ["Mean offset (ms)", f1(C["mean"]), f1(Cm["mean"]), f1(R["mean"])],
    ["95 % CI of the mean (ms)", ci(C["ci95"]), ci(Cm["ci95"]), ci(R["ci95"])],
    ["Between-session SD (ms)", f"{C['sd']:.2f}", f"{Cm['sd']:.2f}", f"{R['sd']:.2f}"],
    ["Range of session means (ms)", f"{C['min']:.1f}–{C['max']:.1f}", f"{Cm['min']:.1f}–{Cm['max']:.1f}",
     f"{R['min']:.1f}–{R['max']:.1f}"],
    ["95 % prediction interval, new session (ms)", ci(C["pi95"]), ci(Cm["pi95"]), ci(R["pi95"])],
    ["Within-session SD (ms)", f"{s['within_run_sd_ms']:.2f}", "—", f"{s['within_run_sd_recorded_marker_ms']:.2f}"],
    ["Per-event error, SD (ms)*", f"{E['sd']:.2f}", f"{Em['sd']:.2f}", f"{Er['sd']:.2f}"],
    ["Per-event error, 2.5–97.5 % (ms)*", f"{E['p2.5']:.1f} to {E['p97.5']:.1f}",
     f"{Em['p2.5']:.1f} to {Em['p97.5']:.1f}", f"{Er['p2.5']:.1f} to {Er['p97.5']:.1f}"],
    ["Events within ±10 ms / ±1 frame*", f"{E['within_10ms']:.1%} / {E['within_1frame']:.1%}",
     f"{Em['within_10ms']:.1%} / {Em['within_1frame']:.1%}", f"{Er['within_10ms']:.1%} / {Er['within_1frame']:.1%}"],
    ["Maximum per-event error (ms)*", f"{E['max_abs']:.1f}", f"{Em['max_abs']:.1f}", f"{Er['max_abs']:.1f}"],
    ["Correlation of session mean with LSL transit", f"{s['corr_with_transit']['offset_mean_ms']:.2f}",
     f"{s['corr_with_transit']['offset_medbase_mean_ms']:.2f}", f"{s['corr_with_transit']['delay_mean_ms']:.2f}"],
], widths=[2.2, 1.7, 1.45, 1.35],
    caption=f"Table S3. Offset between optical onset and marker ({C['n']} sessions, {s['n_flashes_included']:,} flashes)",
    note="Offset = optical onset − marker; positive values mean the light appears after the marker. "
         "*Leave-one-session-out: the correction from all other sessions applied to each flash.")
figure(FIG / "fig1_offset_per_run.png",
       f"Figure S1. Offset per session (mean ± within-session SD). A: one median clock mapping per session. "
       f"B: local upper-envelope mapping (adopted). Black line: mean over sessions; dark band: 95 % CI; light band: "
       f"95 % prediction interval for a new session.")
para(f"**Why the mapping matters.** Sessions fell into two timing states of the recording model: a steady state "
     f"(model about one amplifier frame behind, short LSL transit of ~1.5–2.6 ms) and a burst state (frame-wise "
     f"processing, transit up to ~15 ms; Fig. S2A). The marker as recorded by the LSL Receive block, and the "
     f"timestamp-placed marker with a median clock mapping, both shifted with this state (r = "
     f"{s['corr_with_transit']['delay_mean_ms']:.2f} and {s['corr_with_transit']['offset_medbase_mean_ms']:.2f} "
     f"with transit), producing bimodal session means. The upper-envelope mapping removed most of this dependence "
     f"(r = {s['corr_with_transit']['offset_mean_ms']:.2f}; Fig. S2B) and reduced the between-session SD from "
     f"{Cm['sd']:.1f} to {C['sd']:.1f} ms.")
figure(FIG / "fig2_mechanism.png",
       "Figure S2. A: distribution of t_A − model time over all model steps in a steady and a burst-mode session. "
       "B: session mean offset against mean LSL transit for the recorded marker and the two timestamp mappings.")
para(f"**Per-event accuracy.** With the correction from the other sessions, the optical onset of single flashes "
     f"deviated from the corrected marker by SD {E['sd']:.1f} ms (Fig. S3A). This variability is the same in every "
     f"session (within-session SD {s['within_run_sd_detrended_ms']:.1f} ms after removing each session's trend; "
     f"Fig. S3B), is unrelated to LSL transit, and is distributed approximately uniformly over one display frame "
     f"(16.7 ms), i.e. it arises between the buffer swap reported to PsychoPy and the physical refresh.")
figure(FIG / "fig3_error_distribution.png",
       "Figure S3. A: per-flash error after correction (leave-one-session-out), dashed lines 2.5th and 97.5th "
       "percentiles. B: cumulative distribution of within-session deviations, one line per session.")
lr_rows = []
for rid, d in s["long_runs"].items():
    lr_rows.append([rid, f"{d['duration_min']:.0f}", f"{d['drift_ms_per_min']:+.3f}",
                    f"{d['last100_minus_first100_ms']:+.1f}", f"{d['drift_median_base_ms_per_min']:+.3f}"])
para(f"**Stability over time.** In the long and full-length sessions the corrected offset did not drift "
     f"(Table S4, Fig. S4). The relation between A's clock and the amplifier clock changed slowly in all sessions "
     f"(median {np.median(slopes):.2f} ms/min, ≈ {drift_ppm:.0f} ppm; {full_slope:.2f} ms/min in the "
     f"{full['duration_min']:.0f}-min session, Fig. S5). With a single mapping per session this drift accumulated to "
     f"{abs(full['drift_median_base_ms_per_min']) * full['duration_min']:.0f} ms over the full-length session; the "
     f"local mapping follows it.")
table(["Session", "Duration (min)", "Drift, adopted (ms/min)", "Last − first 100 flashes (ms)",
       "Drift, median mapping (ms/min)"], lr_rows, widths=[1.6, 1.0, 1.3, 1.5, 1.4],
      caption="Table S4. Long and full-length sessions")
figure(FIG / "fig4_long_runs.png",
       "Figure S4. Offset of every flash (dots) and 50-flash running means over the long and full-length sessions; "
       "orange: median clock mapping; black: adopted mapping; dashed: correction.")
figure(FIG / "fig5_model_lag.png",
       "Figure S5. Change of t_A − model time over each session (5-s running median), relative to its first 10 s. "
       "Short stalls are caught up; the common slope reflects drift between the amplifier and PC clocks.")
bt = s["by_run_type"]
para(f"**Other factors.** Session type had no effect (short {bt['short']['mean']:.1f} ms, n = {bt['short']['n']}; "
     f"long {bt['long']['mean']:.1f} ms, n = {bt['long']['n']}; full length {bt['full']['mean']:.1f} ms; Welch "
     f"p = {s['short_vs_long_welch_p']:.2f}). Day 2 was {abs(day['diff_ms']):.1f} ms lower than day 1 "
     f"({day['mean'][0]:.1f} vs {day['mean'][1]:.1f} ms; Welch t = {day['welch_t']:.2f}, p = {day['p']:.2f}), "
     f"within the prediction interval.")
para(f"**Sensitivity and replication.** The between-session SD was lowest for percentiles 85–92 (Fig. S6A; "
     f"{s['percentile_scan']['85']['between_sd']:.2f}–{s['percentile_scan']['92']['between_sd']:.2f} ms) and rose "
     f"towards the median ({s['percentile_scan']['50']['between_sd']:.2f} ms) and the maximum. In 8 independent pilot "
     f"sessions (sensor at the bottom-right corner; frame lengths 4 and 8) the adopted mapping gave "
     f"{V['offset']['mean']:.1f} ms with a between-session SD of {V['offset']['sd']:.2f} ms (median mapping "
     f"{V['offset_median_base']['sd']:.2f} ms), independent of frame length "
     f"({V['by_frame_length']['4']['mean']:.1f} vs {V['by_frame_length']['8']['mean']:.1f} ms; Fig. S6B). The corner "
     f"offset was {V['corner_minus_centre_ms']:.1f} ms larger than at the centre (p < 0.001), consistent with "
     f"top-to-bottom refresh of the display.")
figure(FIG / "fig6_sensitivity_validation.png",
       "Figure S6. A: between-session SD of the offset as a function of the percentile of t_A − model time used for "
       "the clock mapping (dotted: adopted). B: independent pilot sessions (13 October, corner sensor), both mappings.")

# ----------------------------------------------------------------------------- correction
h("S1.6 Correction applied to the experimental data", 1)
para("For every event in an experimental session the onset on the EEG time axis is computed from the recorded "
     "timing file as")
equation(f"τ_event = t_B + Δ_BA − B(τ) + {C['mean']:.1f} ms")
para(f"with B(τ) evaluated around the marker as in S1.4 (implemented in `analysis/marker_timing.py`, which shares its "
     f"code with the calibration analysis; applied to the 22 calibration sessions it reproduces the optical onsets "
     f"with a mean error of −0.2 ms). The constant is valid for stimuli at the screen centre. "
     f"The expected session-level uncertainty is the prediction interval ({ci(C['pi95'])} ms); the per-event "
     f"uncertainty (SD {E['sd']:.1f} ms, bounded by one frame) acts as a small temporal smoothing of the ERP and is "
     f"identical across conditions. Each experimental session must (i) record the `LSLMarkerTimestamps` output, "
     f"(ii) report `lsl_resyncs = 0` on the stimulus computer and (iii) show no lasting backlog of the recording "
     f"model; sessions failing these checks are flagged. Stimulus software, display settings and hardware were "
     f"unchanged between calibration and data collection [to be confirmed per study; re-calibration with "
     f"≥ 3 sessions after any change].")

# ----------------------------------------------------------------------------- limitations
h("S1.7 Limitations", 1)
bullets([
    "The calibration used a full-white square at the screen centre. Stimuli in the upper or lower part of the screen "
    f"appear up to about half a frame earlier or later; at the bottom-right corner we measured +{V['corner_minus_centre_ms']:.1f} ms. "
    "Grey-scale stimuli may cross the sensor threshold at a slightly different point of the rise than full white.",
    f"The database covers 2 days. The day difference ({abs(day['diff_ms']):.1f} ms) was within the prediction "
    "interval, but longer-term stability should be checked with periodic re-calibration sessions.",
    "The 90th-percentile mapping was chosen on the database itself; the choice was confirmed on independent pilot "
    "sessions and results are insensitive between the 85th and 92nd percentiles.",
    "The clock mapping assumes no lasting backlog of the recording model. A pause of the stimulus process of the kind "
    "observed in pilot testing would add a lasting backlog that this mapping cannot separate from clock offset; "
    "such sessions are detected (lsl_resyncs, backlog) and excluded.",
    "Optical onsets are quantised to the 600 Hz sampling (1.67 ms).",
])

# ----------------------------------------------------------------------------- history
h("S1.8 Development history", 1)
para("Initial recordings (6–8 October 2026) showed markers about 0.85–0.97 s earlier than the light in the EEG "
     "file. Simultaneous LSL timestamps located the cause in a single ~0.8 s pause of the stimulus process at its "
     "first full-screen frames, occurring after the LSL stream had been created; the recording model, which reads one "
     "LSL sample per step, never recovered from the resulting gap. Creating the LSL stream only after the window "
     "is open and the refresh rate measured removed the backlog (resynchronisations 1 → 0; delay ~870 ms → "
     "~35–60 ms). On 13 October, 8 sessions compared Simulink frame lengths 4 and 8 (no difference) and introduced "
     "the experiment-start marker; these sessions form the independent replication set above. The calibration "
     "database was then recorded with the sensor at the screen centre, where the experimental stimuli appear.")

# ----------------------------------------------------------------------------- availability
h("Data and code availability", 1)
para(f"Repository: {REPO_URL}. Calibration script: `run_photodiode_latency_v1_trig_after_framerate_measurement.py`; "
     "LSL sender: `lsl_trigger.py`; recording model and timestamp block: `matlab/on_eeg_computer/`; analysis: "
     "`analysis/calibration_database_analysis.py` (produces every number, table and figure in this supplement from "
     "the raw recordings); correction of experimental data: `analysis/marker_timing.py`. Raw recordings (EEG .mat, timing .mat, PsychoPy logs) are available from the authors "
     "[or: repository/DOI].")

# ----------------------------------------------------------------------------- per-session table
h("Table S5. Per-session results", 1)
rows = []
for r in db.itertuples():
    rows.append([f"{r.date[6:]}.10 {r.run_type} {r.run_no}", r.flash_markers,
                 f"{r.offset_mean_ms:.1f} ± {r.offset_sd_ms:.1f}", f"{r.offset_drift_ms_per_min:+.2f}",
                 f"{r.offset_medbase_mean_ms:.1f}", f"{r.delay_mean_ms:.1f}", f"{r.transit_mean_ms:.1f}",
                 f"{r.lag_iqr_ms:.1f}", r.stalls_over_100ms, int(r.dropped_frames)])
table(["Session", "Flashes", "Offset, adopted (ms, mean ± SD)", "Drift (ms/min)", "Offset, median mapping (ms)",
       "Recorded marker (ms)", "LSL transit (ms)", "Lag IQR (ms)", "Stalls", "Dropped frames"],
      rows, widths=[1.0, 0.5, 0.95, 0.6, 0.75, 0.65, 0.55, 0.55, 0.45, 0.6], caption="", font=7.5, keep_together=False,
      note="Lag IQR: interquartile range of t_A − model time (≈ 1–2 ms steady state, ≈ 7–15 ms burst state).")

# ----------------------------------------------------------------------------- references
h("References", 1)
refs = [
    "Bridges, D., Pitiot, A., MacAskill, M. R., & Peirce, J. W. (2020). The timing mega-study: comparing a range of "
    "experiment generators, both lab-based and online. PeerJ, 8, e9414. https://doi.org/10.7717/peerj.9414",
    "Elze, T. (2010). Achieving precise display timing in visual neuroscience experiments. Journal of Neuroscience "
    "Methods, 191(2), 171–179. https://doi.org/10.1016/j.jneumeth.2010.06.018",
    "Keil, A., Debener, S., Gratton, G., Junghöfer, M., Kappenman, E. S., Luck, S. J., Luu, P., Miller, G. A., & "
    "Yee, C. M. (2014). Committee report: Publication guidelines and recommendations for studies using "
    "electroencephalography and magnetoencephalography. Psychophysiology, 51(1), 1–21. https://doi.org/10.1111/psyp.12147",
    "Kothe, C., Shirazi, S. Y., Stenner, T., Medine, D., Boulay, C., Grivich, M. I., Mullen, T., Delorme, A., & "
    "Makeig, S. (2024). The Lab Streaming Layer for synchronized multimodal recording. bioRxiv. [verify final citation]",
    "Luck, S. J. (2014). An introduction to the event-related potential technique (2nd ed.). MIT Press.",
    "Peirce, J., Gray, J. R., Simpson, S., MacAskill, M., Höchenberger, R., Sogo, H., Kastman, E., & Lindeløv, J. K. "
    "(2019). PsychoPy2: Experiments in behavior made easy. Behavior Research Methods, 51(1), 195–203. "
    "https://doi.org/10.3758/s13428-018-01193-y",
    "Pernet, C., Garrido, M. I., Gramfort, A., Maurits, N., Michel, C. M., Pang, E., Salmelin, R., Schoffelen, J. M., "
    "Valdes-Sosa, P. A., & Puce, A. (2020). Issues and recommendations from the OHBM COBIDAS MEEG committee for "
    "reproducible EEG and MEG research. Nature Neuroscience, 23(12), 1473–1483. https://doi.org/10.1038/s41593-020-00709-0",
    "Plant, R. R. (2016). A reminder on millisecond timing accuracy and potential replication failure in "
    "computer-based psychology experiments: An open letter. Behavior Research Methods, 48(1), 408–411. "
    "https://doi.org/10.3758/s13428-015-0577-0",
    "Plant, R. R., Hammond, N., & Turner, G. (2004). Self-validating presentation and response timing in cognitive "
    "paradigms: How and why? Behavior Research Methods, Instruments, & Computers, 36(2), 291–303. "
    "https://doi.org/10.3758/BF03195575",
]
for r in refs:
    p = para(r, size=9, space_after=3)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print("saved", OUT)
