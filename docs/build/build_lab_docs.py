"""Build docs/lab_guide_marker_timing_2026-10.docx and docs/todo_experiment_and_analysis_2026-10.docx."""
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(r"C:\Users\aviva-Lab\Desktop\vMMR_EEG_PsychoPy_experiment")
REPO_URL = "https://github.com/DanielAtad/vMMR_EEG_PsychoPy_experiment"
EXP = "run_vMMR_experiment_v1_trig_after_framerate_measurement.py"


class Doc:
    def __init__(self):
        self.d = docx.Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        for sec in self.d.sections:
            sec.left_margin = sec.right_margin = Inches(0.9)
            sec.top_margin = sec.bottom_margin = Inches(0.8)

    def runs(self, p, text, bold=False, size=None):
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 10.5) - 1)
            else:
                r = p.add_run(part); r.bold = bold
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
        return p

    def p(self, text="", bold=False, size=None, after=6, italic=False):
        p = self.d.add_paragraph()
        self.runs(p, text, bold, size)
        if italic:
            for r in p.runs:
                r.italic = True
        p.paragraph_format.space_after = Pt(after)
        return p

    def h(self, text, level):
        self.d.add_heading(text, level=level)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            p = self.d.add_paragraph(style=style)
            self.runs(p, it)
            p.paragraph_format.space_after = Pt(2)

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            p = self.d.add_paragraph()
            r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Inches(0.3)
        self.p("", after=4)

    def table(self, header, rows, widths, caption=None, font=9):
        if caption:
            self.p(caption, bold=True, size=9.5, after=3).paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, x in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""
            self.runs(c.paragraphs[0], x, bold=True, size=font)
            tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E8EEF7")
            tcPr.append(shd)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""; self.runs(cells[i].paragraphs[0], str(v), size=font)
        for ri, row in enumerate(t.rows):
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
        self.p("", after=4)

    def save(self, path):
        self.d.save(path)
        print("saved", path)


# =============================================================================
# 1. Lab guide
# =============================================================================
g = Doc()
g.d.add_heading("vMMR EEG: marker timing – how to record and how to correct", 0)
g.p(f"Lab guide, 15 October 2026. Repository: {REPO_URL} (branch `photodiode-latency-calibration`). "
    "Methods and results: Supplementary Material S1 (`docs/supplement_marker_timing_calibration_2026-10-15.docx`).",
    size=9, after=10)

g.h("1. In one paragraph", 1)
g.p("Event markers are sent by the PsychoPy computer (B) over LSL and recorded with the EEG on the g.HIamp/Simulink "
    "computer (A). The marker channel inside the EEG file is **not** accurate enough for ERP timing (it varies by "
    "about 9 ms between sessions). Instead, every event is placed by its LSL send timestamp, which a second LSL "
    "block saves in a separate timing file, and a calibrated constant is added: **onset = timestamp-placed marker "
    "+ 62.4 ms**. This was measured with an optical sensor in 22 sessions (between-session SD 1.4 ms; single events "
    "SD 5.7 ms, all within one screen frame). The constant is valid only for the configuration below and for "
    "stimuli at the screen centre.")

g.h("2. Computer B – stimulus PC (PsychoPy)", 1)
g.table(["Item", "Required setting"], [
    ["Display", "Dell P2426H, 1920×1080. Windows: Settings → System → Display → Advanced display → refresh rate "
                "**59.94 Hz** (not 60, not 120). Check before each session day: `.venv\\Scripts\\python "
                "check_display_refresh_rate.py` must report ≈ 59.94 Hz and 100 % of intervals in 14–19 ms."],
    ["Software", "Project `.venv`: Python 3.10.22, PsychoPy 2026.2.4, pylsl 1.18.5. Do not update without "
                 "re-calibrating."],
    ["Network", "`lsl_api.cfg` in the project folder with `KnownPeers = {192.168.50.1}` (computer A)."],
    ["Experiment script", f"`{EXP}` (until the updated version from the to-do document exists). Dialog: "
                          "`expected_refresh_hz` 59.94, `lsl_keepalive_hz` 600, `lsl_nominal_srate` 600, "
                          "`send_LSL_triggers` on, `photodiode_square` **off** for real participants (the current "
                          "default is on)."],
    ["Files needed", f"`{EXP}`, `lsl_trigger.py`, `timing_utils.py`, `conditions/`, `stimuli/`, "
                     "`check_display_refresh_rate.py`, `lsl_api.cfg`."],
    ["After every session", "Open the newest `data/*_run_info.txt`: `lsl_resyncs: 0` is **required**. Keep all "
                            "outputs (`_events.csv`, `_trials.csv`, `_run_info.txt`, `_frame_intervals.csv`, `.log`)."],
], widths=[1.4, 5.3])

g.h("3. Computer A – recording PC (g.HIamp, Simulink)", 1)
g.table(["Item", "Required setting"], [
    ["Software", "MATLAB/Simulink R2020a with g.HIsys; **liblsl-Matlab on the MATLAB path** (`lsl_loadlib` must work)."],
    ["Model", "`matlab/on_eeg_computer/HIamp_record_LSL_DD_vMMR.slx` with `LSLMarkerTimestamps.m` in the same "
              "folder (or on the path). Do not use the older `.slx` files in the repository root."],
    ["g.HIamp block", "600 Hz, **frame length 8** (block configuration dialog)."],
    ["Solver / model settings", "Fixed-step, discrete, step 1/600, single-tasking (“treat each discrete rate as a "
                                "separate task” unchecked), simulation pacing off, multithreaded co-simulation on "
                                "(model on, g.HIamp block auto), workspace logging on."],
    ["LSL blocks", "`LSL MARKERS FROM EXPERIMENT` (g.HIsys LSL Receive): stream `experiment_markers`, sampleRate "
                   "600. `LSL MARKER TIMESTAMPS` (MATLAB System, interpreted execution): StreamName "
                   "`experiment_markers`, ResolveTimeout 5, CorrectionInterval 1."],
    ["Output files", "Two To File blocks, same base name: `<name>.mat` (variable `EEG_data`) and "
                     "`<name>_LSLtiming.mat` (variable `LSL_timing`, Timeseries). Rename both before every session; "
                     "**the timing file is required for the correction.**"],
    ["Stop", "The model stops on marker 255. The current experiment script ends with 99, so stop the model "
             "manually after the end screen until the script is updated."],
], widths=[1.4, 5.3])

g.h("4. Session order", 1)
g.bullets([
    "B: start the experiment script, fill in the dialog. The window opens and the refresh rate is measured.",
    "B: the LSL stream is created and the screen asks to start the recording model.",
    "A: set both output file names, then start the Simulink model (the stream must already exist: the timestamp "
    "block looks for it at model start, for 5 s).",
    "B: the script continues when Simulink is connected; run the experiment.",
    "End: stop the model (see above), check `lsl_resyncs: 0` on B, copy the two .mat files from A and all outputs "
    "from B to the study folder.",
], style="List Number")

g.h("5. Applying the correction (analysis)", 1)
g.p("`analysis/marker_timing.py` does the whole correction for one recording. It needs `numpy`, `pandas` and "
    "`h5py` (the .mat files are MATLAB v7.3/HDF5).")
g.code("""
import sys; sys.path.insert(0, "analysis")    # repository analysis folder
from marker_timing import correct_events, to_mne_events

# the ..._LSLtiming.mat file must be next to the EEG file
events, qc = correct_events("main_..._sub01.mat",
                            run_info="sub01_..._run_info.txt")
if not qc["passed"]:
    print("CHECK THIS SESSION:", qc["problems"])
events[["code", "onset_s", "sample"]]         # corrected time (s) and EEG row
mne_events = to_mne_events(events, codes=[11, 12, 21, 22, 31, 32, 41, 42])
""")
g.p("Or from the command line (writes `<recording>_events_corrected.csv`):")
g.code(r""".venv\Scripts\python analysis\marker_timing.py main_..._sub01.mat ^
    --run-info sub01_..._run_info.txt""")
g.table(["Column / field", "Meaning"], [
    ["`code`", "marker code (11–44 face events, 9 start, 99 end, …)"],
    ["`onset_s`", "corrected onset on the EEG time axis (column 0 of `EEG_data`), in seconds"],
    ["`sample`", "row of `EEG_data` closest to the corrected onset (use as the event sample)"],
    ["`recorded_s`, `shift_vs_recorded_ms`", "where the marker channel had the event, and the shift applied"],
    ["`ts_B`", "LSL send timestamp on computer B (matches the PsychoPy LSL timestamp of the event)"],
    ["`qc`", "`passed`, `problems`, `lsl_resyncs`, `lasting_backlog_ms`, `stalls_over_100ms`, clock drift"],
], widths=[2.0, 4.7])
g.bullets([
    "**Matching events to the PsychoPy events file:** face events (codes 11–44) appear in the same order as the rows "
    "of `*_events.csv`; check that the code sequences are identical. After the planned script update the events file "
    "will also contain each event's LSL timestamp (`ts_B`), which allows exact matching.",
    "**Practice trials** send the same codes; drop them using the events file (`phase`).",
    "**Session check:** use only sessions with `qc[\"passed\"]` (lsl_resyncs 0, no lasting backlog of the model).",
    "**Stimulus position:** 62.4 ms is for the screen centre (faces). The prime word (260 px above the centre) "
    "appears about 4 ms earlier because the screen is drawn top to bottom (estimate, not measured).",
    "**Never** add 62.4 ms to the marker channel of `EEG_data`; the constant belongs to the timestamp method only.",
])

g.h("6. When to re-calibrate", 1)
g.p("After any change of monitor, graphics driver, Windows display settings, PsychoPy/Python packages, "
    "`lsl_trigger.py`, the Simulink model or g.HIamp settings: record at least 3 calibration sessions with "
    "`run_photodiode_latency_v1_trig_after_framerate_measurement.py` (optical sensor at the screen centre) and "
    "analyse them with `analysis/calibration_database_analysis.py`. If they fall outside 59.4–65.3 ms, a new "
    "constant is needed.")

g.h("7. Where things are", 1)
g.table(["What", "Where"], [
    ["Correction function", "`analysis/marker_timing.py`"],
    ["Calibration analysis", "`analysis/calibration_database_analysis.py`"],
    ["Calibration data", "Google Drive: `final database/` (`EEG_Computer_A/`, `Computer_B/`, `run_manifest.csv`)"],
    ["Recording model (A)", "`matlab/on_eeg_computer/`"],
    ["Methods supplement", "`docs/supplement_marker_timing_calibration_2026-10-15.docx` / `.pdf`"],
    ["Open work", "`docs/todo_experiment_and_analysis_2026-10.docx`"],
], widths=[1.8, 4.9])
g.save(REPO / "docs" / "lab_guide_marker_timing_2026-10.docx")

# =============================================================================
# 2. To-do
# =============================================================================
t = Doc()
t.d.add_heading("To do: experiment script update and vMMR analysis script", 0)
t.p(f"15 October 2026. Repository: {REPO_URL}. Line numbers refer to `{EXP}` at commit `34fc744`.", size=9, after=10)

t.h("Part 1 – Experiment script", 1)
t.p(f"Create `run_vMMR_experiment_v2.py` from `{EXP}` (keep v1 unchanged, as the README asks). Items 1–5 are "
    "needed before data collection; 6–8 are recommended.")
t.table(["#", "Change", "Why", "Where (v1)"], [
    ["1", "End the session with **255** (send 99 = end, then 255 = stop)",
     "The recording model stops only on 255; with 99 it keeps recording until stopped by hand.",
     "`finish(final_code=99)` :1364, `set(99)` :1367"],
    ["2", "**Preload images and read the trial tables before creating the LSL stream**",
     "Loading textures while the stream is live can pause PsychoPy; a pause > 0.1 s is a resync and shifts "
     "everything after it (the 0.8 s stall found in pilot testing). Same rule as for the window/refresh measurement.",
     "stream created :1097; CSVs :1236–1238; `preload_images` :1254"],
    ["3", "**Write each face event's LSL timestamp to the events file** (column `lsl_event_timestamp`)",
     "Already captured (`onset[\"lsl\"]`, :552) but not saved. Lets the analysis match each events-file row to its "
     "corrected event exactly instead of relying on order.",
     "`make_event_writer` fields :387–390; `present_face_event` return; `event_writer.writerow` :898"],
    ["4", "**Experiment-start marker 1**, flip-aligned, when SPACE ends the instructions",
     "Same anchor as the calibration scripts and the AGENTS.md map; marks the start after the model is "
     "running. Keep 9 (script start, not flip-aligned) or drop it.",
     "`trigger.set(9)` :1283; instructions :1285"],
    ["5", "**Photodiode square off by default** (`photodiode_square`); photodiode test trigger not 99",
     "The square must not appear for participants. `PHOTODIODE_TEST_TRIGGER = 99` reuses the end code (and would "
     "stop an old model).", "dialog :954; :110"],
    ["6", "Write to `run_info`: git commit, file hashes, software versions, Windows refresh setting",
     "Traceability for the methods section; the 120 Hz incident of 14 October was found only by chance.",
     "run_info block :1150–1190 (as in `run_photodiode_latency_database.py`)"],
    ["7", "Optional: prime-onset (51–62) and fixation (80) markers via `win.callOnFlip`",
     "Listed in AGENTS.md but not sent; useful for prime-locked analyses and epoch checks. Not needed for the vMMR.",
     "`run_trial` :828"],
    ["8", "Optional: show `lsl_resyncs` on the end screen",
     "Immediate check by the experimenter instead of opening run_info.", "end screen :1334"],
], widths=[0.25, 1.75, 2.75, 1.95], font=8.5)
t.p("**Validation after the update** (before the first participant):")
t.bullets([
    "Dry run with the AGENTS.md checklist (dialog, practice and main trials, ESC, Hebrew, preloading, one events row "
    "per face, targets marked, square only on face-on frames when enabled).",
    "1–3 full-length sessions with the photodiode square on face-on frames and the sensor at the **bottom-right "
    "corner** (the existing square position, so faces are not covered). Apply `correct_events`, compare with the "
    "optical onsets per face event: expected +2.0 ms (corner, from the 13 October sessions), SD ≈ 6 ms; also "
    "`lsl_resyncs` 0, model stopped by 255, all codes 11–44 present, events-file timestamps matched. Then set "
    "the square off for real sessions.",
])

t.h("Part 2 – vMMR analysis script", 1)
t.p("Nothing exists yet besides two notebooks that only load a June test file. Proposed structure (Python/MNE, "
    "or the lab's EEGLAB/FieldTrip pipeline with the same steps; decide first):")
t.table(["Step", "Content", "Open decisions"], [
    ["1 Load", "EEG .mat (`EEG_data`: col 0 time, EEG channels, optical sensor, marker) → MNE Raw with channel "
               "names, units and montage", "Channel order/names and units of the g.HIamp configuration; cap layout"],
    ["2 Events", "`marker_timing.correct_events` (+ QC); keep the main phase; match to `*_events.csv` "
                 "(by LSL timestamp after Part 1 item 3, until then by code order); attach metadata (prime, identity, "
                 "role, `isAnalysisStandard`, `isAnalysisDeviant`, target trials)", "Session exclusion if QC fails"],
    ["3 Preprocessing", "Filtering, 50 Hz line noise, bad channels, re-reference, ICA for eye artifacts",
     "Filter cut-offs, reference, ICA settings – follow the original paradigm or preregistration"],
    ["4 Epochs", "−100 to +500 ms around corrected face onset, baseline −100–0 ms (README); 3rd standard vs "
                 "deviant, non-target trials only; artifact rejection", "Rejection thresholds / autoreject"],
    ["5 ERPs", "Per participant: 2 × 2 (prime DEATH/NEGATIVE × identity SELF/OTHER) standard and deviant ERPs; "
               "vMMR = deviant − standard", "ROIs and time windows"],
    ["6 Statistics", "Group level: repeated-measures analysis of window means and/or cluster-based permutation tests",
     "Primary test as preregistered"],
    ["7 Behaviour", "Target hit rate, false alarms, RT from `*_trials.csv`; participant exclusion criteria", "Criteria"],
    ["8 Output", "Per-participant QC report (n epochs per condition, rejected epochs, timing QC), group figures, "
                 "tables; one config file; tests on a short real recording", "BIDS export?"],
], widths=[0.9, 3.6, 2.2], font=8.5)
t.p("Timing in the analysis: use only the corrected onsets; the per-event timing SD (5.7 ms, within one frame) is the "
    "same in all conditions and acts as slight smoothing. Prime-locked analyses should note that the prime appears "
    "about 4 ms before the centre (estimate).", after=6)
t.save(REPO / "docs" / "todo_experiment_and_analysis_2026-10.docx")
